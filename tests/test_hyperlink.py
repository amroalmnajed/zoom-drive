import sys
import docx

def add_hyperlink(paragraph, url, text):
    """
    A function that places a hyperlink within a paragraph object.
    """
    # This gets access to the document.xml.rels file and gets a new relation id value
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    # Create the w:hyperlink tag and add needed values
    hyperlink = docx.oxml.shared.OxmlElement('w:hyperlink')
    hyperlink.set(docx.oxml.shared.qn('r:id'), r_id)

    # Create a w:r element
    new_run = docx.oxml.shared.OxmlElement('w:r')

    # Create a new w:rPr element
    rPr = docx.oxml.shared.OxmlElement('w:rPr')

    # Join all the xml elements together add add the required text to the w:r element
    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)

    paragraph._p.append(hyperlink)

    return hyperlink

def modify_docx(path, out_path):
    doc = docx.Document(path)
    table = doc.tables[0]
    # Let's modify row 1 for testing
    cell = table.rows[1].cells[4] # 'Link'
    cell.text = '' # Clear existing text
    p = cell.paragraphs[0]
    add_hyperlink(p, 'https://google.com', 'Link')
    doc.save(out_path)

modify_docx(r'C:\Users\lenovo\Desktop\project\Main_Revision .docx', 'test_out.docx')
