import io
import sys
from datetime import datetime
from pathlib import Path
import docx

# Add project root to sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx_updater import (
    link_session_in_document,
    revert_link_in_document,
    _cell_has_link,
)

URL = "https://drive.google.com/file/d/ABC/view"
OCT10 = datetime(2026, 10, 10, 9, 0)


def build_main():
    d = docx.Document()
    d.add_paragraph("Main Revision Records")
    t = d.add_table(rows=1, cols=5)
    for c, h in zip(t.rows[0].cells, ["Session #", "Date", "Day", "Content", "Link"]):
        c.text = h
    rows = [
        (40, "4/10/2026", "Sunday"),
        (41, "5/10/2026", "Monday"),
        (42, "6/10/2026", "Tuesday"),
        (43, "10/10/2026", "Saturday"),
        (44, "11/10/2026", "Sunday"),
        (45, "12/10/2026", "Monday"),
        (47, "18/10/2026", ""),
    ]
    for s, dt, day in rows:
        cells = t.add_row().cells
        cells[0].text, cells[1].text, cells[2].text, cells[4].text = str(s), dt, day, "Link"
    return d, t


def build_crash():
    d = docx.Document()
    d.add_paragraph("Crash Course Records")
    t1 = d.add_table(rows=1, cols=5)
    for c, h in zip(t1.rows[0].cells, ["Session #", "Date", "Day", "Content", "Link"]):
        c.text = h
    for s, dt, day in [(21, "10-3", "Saturday"), (25, "10-7", "Wednesday")]:
        r = t1.add_row().cells
        r[0].text, r[1].text, r[2].text, r[4].text = str(s), dt, day, "link"
    d.add_page_break()
    d.add_paragraph("Crash Course Records")
    t2 = d.add_table(rows=0, cols=5)  # NO header row
    for s, dt, day in [(26, "10-10", "Saturday"), (27, "10-11", "Sunday"), (28, "10-12", "Monday")]:
        r = t2.add_row().cells
        r[0].text, r[1].text, r[2].text, r[4].text = str(s), dt, day, "link"
    return d, t1, t2


def roundtrip(d):
    b = io.BytesIO()
    d.save(b)
    b.seek(0)
    return docx.Document(b)


def cell_for(d, ti, ri, ci):
    return d.tables[ti].rows[ri].cells[ci]


def test_main_s43():
    """Main S43 links correctly without matching row 40."""
    d, t = build_main()
    for ri, s in ((1, 40), (2, 41), (3, 42)):
        r = link_session_in_document(d, s, f"https://x/{s}", None)
        assert r.status == "updated"
    r = link_session_in_document(d, 43, URL, OCT10)
    assert r.status == "updated", r
    d = roundtrip(d)
    assert _cell_has_link(cell_for(d, 0, 4, 4)) and cell_for(d, 0, 4, 4).text == "Link"
    assert not _cell_has_link(cell_for(d, 0, 5, 4))  # S44 untouched


def test_already_linked():
    """Already linked row reports cleanly as status."""
    d, t = build_main()
    link_session_in_document(d, 43, URL, OCT10)
    d = roundtrip(d)
    r = link_session_in_document(d, 43, URL + "2", OCT10)
    assert r.status == "already_linked", r


def test_crash_s26_second_table():
    """Crash S26 is found in the second header-less table."""
    d, t1, t2 = build_crash()
    r = link_session_in_document(d, 26, URL, OCT10)
    assert r.status == "updated" and "table 2, row 1" in r.detail, r
    d = roundtrip(d)
    assert _cell_has_link(cell_for(d, 1, 0, 4)) and not _cell_has_link(cell_for(d, 1, 1, 4))


def test_unknown_session():
    """Unknown session returns not_found status."""
    d, t = build_main()
    r = link_session_in_document(d, 99, URL, OCT10)
    assert r.status == "not_found"


def test_date_fallback():
    """Date fallback works for blank session cell, but not wrong session number."""
    d, t = build_main()
    t.rows[7].cells[0].text = ""  # blank session cell, date 18/10/2026
    r = link_session_in_document(d, 50, URL, datetime(2026, 10, 18, 9, 0))
    assert r.status == "updated", r

    # Date does not override a different session number
    d, t = build_main()
    r = link_session_in_document(d, 99, URL, datetime(2026, 10, 18, 9, 0))
    assert r.status == "not_found", r


def test_link_styling():
    """Google blue #1155CC + underline and font preserved."""
    d, t = build_main()
    cell = t.rows[4].cells[4]
    cell.paragraphs[0].runs[0].font.name = "Arial"
    cell.paragraphs[0].runs[0].font.size = docx.shared.Pt(11)
    link_session_in_document(d, 43, URL, OCT10)
    d = roundtrip(d)
    xml = cell_for(d, 0, 4, 4)._tc.xml
    assert 'w:val="1155CC"' in xml and '<w:u w:val="single"/>' in xml and 'Arial' in xml and 'w:hyperlink' in xml


def test_revert_link():
    """Undo restores cell, drops relationship, 2nd undo is a no-op."""
    d, t = build_main()
    r = link_session_in_document(d, 43, URL, OCT10)
    d = roundtrip(d)
    assert revert_link_in_document(d, URL, r.original_cell_xml) == "reverted"
    d = roundtrip(d)
    assert not _cell_has_link(cell_for(d, 0, 4, 4)) and cell_for(d, 0, 4, 4).text == "Link"
    assert URL not in {rel.target_ref for rel in d.part.rels.values()}
    assert revert_link_in_document(d, URL, r.original_cell_xml) == "not_found"
