import io, sys, types, importlib.abc, importlib.machinery, os, json, tempfile
from unittest.mock import MagicMock
from datetime import datetime, timezone
sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parents[1]))

# --- stub libs that aren't installed in this sandbox ---
class HttpError(Exception):
    def __init__(self, status): self.resp = types.SimpleNamespace(status=status)
class _Stub(importlib.abc.MetaPathFinder, importlib.abc.Loader):
    names = ("googleapiclient", "google", "dotenv")
    def find_spec(self, name, path, target=None):
        if name.split(".")[0] in self.names:
            return importlib.machinery.ModuleSpec(name, self, is_package=True)
    def create_module(self, spec):
        m = MagicMock(); m.__path__ = []; m.__spec__ = spec; m.__name__ = spec.name
        if spec.name == "googleapiclient.errors": m.HttpError = HttpError
        if spec.name == "googleapiclient.http":
            class Down:
                def __init__(self, fh, req): self.fh, self.req = fh, req
                def next_chunk(self): self.fh.write(self.req.data); return None, True
            m.MediaIoBaseDownload = Down
            class Up:
                def __init__(self, fh, mimetype=None, resumable=False): self.data = fh.read()
            m.MediaIoBaseUpload = Up
        return m
    def exec_module(self, m): pass
try:
    import googleapiclient  # noqa: F401
except ImportError:
    sys.meta_path.insert(0, _Stub())

import docx
import config
tmp = tempfile.mkdtemp()
config.STATE_FILE = __import__("pathlib").Path(tmp) / "processed.json"
config.UNDO_FILE = __import__("pathlib").Path(tmp) / "undo.json"
config.UNDO_DELETE_PERMANENTLY = False  # Use trashed=True instead of delete for test
import googleapiclient.http as ghttp
import main, docx_updater
from undo_store import UndoStore
from state import StateStore

class _Down:
    def __init__(self, fh, req):
        self.fh, self.req = fh, req
    def next_chunk(self):
        data = getattr(self.req, "data", b"")
        if data:
            self.fh.write(data)
        return None, True

ghttp.MediaIoBaseDownload = _Down

# --- fake Drive holding one docx and one video ---
def make_docx():
    d = docx.Document(); t = d.add_table(rows=1, cols=5)
    for c, h in zip(t.rows[0].cells, ["Session #","Date","Day","Content","Link"]): c.text = h
    for s, dt in ((42,"6/10/2026"),(43,"10/10/2026")):
        r = t.add_row().cells; r[0].text, r[1].text, r[4].text = str(s), dt, "Link"
    b = io.BytesIO(); d.save(b); return b.getvalue()

class FakeDrive:
    def __init__(self): self.docx = make_docx(); self.video_trashed = False; self.updates = []
    def files(self): return self
    def list(self, **kw): return types.SimpleNamespace(execute=lambda: {"files": [{"id": "DOCX1", "name": "Main_Revision.docx"}]})
    def get(self, fileId, **kw):
        return types.SimpleNamespace(execute=lambda: {
            "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        })
    def get_media(self, fileId, **kw): return types.SimpleNamespace(data=self.docx)
    def update(self, fileId, media_body=None, body=None, **kw):
        def _media_bytes(media):
            data = getattr(media, "data", None)
            if data is not None:
                return data
            fd = getattr(media, "_fd", None)
            if fd is not None:
                fd.seek(0)
                return fd.read()
            return media.getbytes()
        def ex():
            if media_body is not None: self.docx = _media_bytes(media_body)
            if body and body.get("trashed"): self.video_trashed = True
            self.updates.append({"fileId": fileId, "body": body, "media_body": bool(media_body)})
            return {}
        return types.SimpleNamespace(execute=ex)
    def delete(self, fileId, **kw):
        self.video_trashed = True
        return types.SimpleNamespace(execute=lambda: None)

class FakeBot:
    def __init__(self): self.sent = []
    def send_message(self, text, reply_markup=None): self.sent.append((text, reply_markup)); return 1

drive = FakeDrive()
rec = types.SimpleNamespace(local_start=lambda: datetime(2026,10,10,9,0), key="rec-1")
URL = "https://drive.google.com/file/d/VID/view"

info = docx_updater.update_docx_in_drive(drive, "F", rec, URL, 43)
assert info["status"] == "updated", info
assert docx.Document(io.BytesIO(drive.docx)).tables[0].rows[2].cells[4]._tc.xml.count("1155CC") == 1
print("drive wrapper: updated + uploaded OK")
assert docx_updater.update_docx_in_drive(drive, "F", rec, URL+"2", 43)["status"] == "already_linked"

# register undo entry + state exactly as the pipeline does
state = StateStore(config.STATE_FILE); state.mark("rec-1", "uploaded", topic="t")
store = UndoStore(config.UNDO_FILE)
eid = store.add({"rec_key":"rec-1","topic":"t","folder_key":"1","folder_name":"Main-Revision","session":43,
                 "video_name":"GMT20261010_Main_S43.mp4","video_id":"VID","video_uploaded":True,"link":URL,
                 "docx_id":"DOCX1","docx_name":"Main_Revision.docx","docx_updated":True,
                 "original_cell_xml": info["original_cell_xml"]})

main.get_drive_service = lambda folder: drive
bot = FakeBot()
main._handle_undo_list(bot);  assert eid in json.dumps(bot.sent[-1][1]); print("/undo list OK")
main._handle_undo_ask(bot, eid); assert "undoyes:" in json.dumps(bot.sent[-1][1]); print("confirm prompt OK")
main._handle_undo_yes(bot, eid)
# print(bot.sent[-1][0])  # Skip print to avoid Unicode encoding issues on Windows
assert drive.video_trashed
d2 = docx.Document(io.BytesIO(drive.docx))
assert "w:hyperlink" not in d2.tables[0].rows[2].cells[4]._tc.xml and d2.tables[0].rows[2].cells[4].text == "Link"
assert not StateStore(config.STATE_FILE).should_hide("rec-1")      # offered again
assert UndoStore(config.UNDO_FILE).get(eid)["undone"]
main._handle_undo_yes(bot, eid); assert "already undone" in bot.sent[-1][0]
main._handle_undo_list(bot); assert "Nothing to undo" in bot.sent[-1][0]
print("ALL FLOW TESTS PASSED")
